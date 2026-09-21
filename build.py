import csv, unicodedata, re, collections as C
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

BASE = Path(__file__).parent
REF = BASE / "istat"  # elenchi Istat di regioni, province e comuni
SOGLIA = 30.0

def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = s.replace("'", " ").replace("-", " ").replace("’", " ")
    return re.sub(r"\s+", " ", s).strip()

def load_ref(fn):
    return [tuple(r.values()) for r in csv.DictReader(open(REF / fn, encoding="utf-8"))]

# --- Regioni
REG = {norm(n): (n, c) for n, c in load_ref("regioni_istat.csv")}
REG_ALIAS = {"friuli venezia g.": "friuli venezia giulia"}

# --- Province: MIM usa ancora l'assetto sardo pre-2025 → codici Istat in vigore fino al 2025
PROV = {norm(n): (n, c) for n, c in load_ref("province_istat.csv")}
PROV.update({
    "sassari": ("Sassari", "090"), "nuoro": ("Nuoro", "091"), "cagliari": ("Cagliari", "292"),
    "oristano": ("Oristano", "095"), "sud sardegna": ("Sud Sardegna", "111"),
    "reggio emilia": ("Reggio nell'Emilia", "035"), "forli cesena": ("Forlì-Cesena", "040"),
})

# --- Comuni (con omonimi)
COM = C.defaultdict(list)
for n, c in load_ref("comuni_istat.csv"):
    COM[norm(n)].append((n, c))

SARDEGNA = {"090": {"112", "113"}, "091": {"114", "116"}, "292": {"118"}, "095": {"115"}, "111": {"114", "116", "117", "118", "119"}}
COM_ALIAS = {"castegnero": "castegnero nanto", "nanto": "castegnero nanto", "montemagno": "montemagno monferrato",
             "murisengo": "murisengo monferrato", "popoli": "popoli terme", "puegnago sul garda": "puegnago del garda",
             "san dorligo della valle dolina": "san dorligo della valle"}

def prefissi(cod_prov):
    if cod_prov in SARDEGNA: return SARDEGNA[cod_prov]
    return {"0" + cod_prov[1:]} if cod_prov[0] == "2" else {cod_prov}

def comune(nome, cod_prov):
    k = COM_ALIAS.get(norm(nome), norm(nome))
    pre = prefissi(cod_prov)
    cands = [x for x in COM.get(k, []) if x[1][:3] in pre]
    if not cands:  # nomi troncati a 30 caratteri nell'anagrafica MIM
        cands = [x for kk, v in COM.items() if len(k) >= 25 and kk.startswith(k) for x in v if x[1][:3] in pre]
    return cands[0] if len(cands) == 1 else (nome.title(), "")

# --- Anagrafica
ana = {r["CODICESCUOLA"]: r for r in csv.DictReader(open(BASE / "SCUANAGRAFESTAT20242520250831.csv", encoding="utf-8-sig"))}

ORDINI = {"SCUOLA PRIMARIA": "Primaria", "SCUOLA SECONDARIA I GRADO": "Secondaria di primo grado"}
righe = [r for r in csv.DictReader(open(BASE / "ALUITASTRACITSTA20242520250831.csv", encoding="utf-8-sig"))
         if r["ORDINESCUOLA"] in ORDINI]

non_match = []

def build(campo_str, etichetta, out_name):
    scuole = C.OrderedDict()
    for r in righe:
        k = (r["CODICESCUOLA"], ORDINI[r["ORDINESCUOLA"]])
        d = scuole.setdefault(k, {"anni": {}, "tot": 0, "str": 0})
        a, s = int(r["ALUNNI"]), int(r[campo_str])
        d["anni"][int(r["ANNOCORSO"])] = (a, s)
        d["tot"] += a; d["str"] += s

    prov = C.defaultdict(lambda: C.Counter())
    elenco = []
    for (cod, ordine), d in scuole.items():
        an = ana[cod]
        reg = REG[REG_ALIAS.get(norm(an["REGIONE"]), norm(an["REGIONE"]))]
        pv = PROV[norm(an["PROVINCIA"])]
        com = comune(an["DESCRIZIONECOMUNE"], pv[1])
        if not com[1]:
            non_match.append(an["DESCRIZIONECOMUNE"])
        pct_anni = {y: round(s / a * 100, 1) if a else None for y, (a, s) in d["anni"].items()}
        anni_over = sorted(y for y, (a, s) in d["anni"].items() if a and s / a * 100 > SOGLIA)
        pct_tot = d["str"] / d["tot"] * 100 if d["tot"] else None
        tot_over = pct_tot is not None and pct_tot > SOGLIA
        key = (reg, pv)
        c = prov[key]
        c[ordine + "|n"] += 1
        c[ordine + "|anno"] += bool(anni_over)
        c[ordine + "|tot"] += tot_over
        if anni_over or tot_over:
            elenco.append([
                cod, an["DENOMINAZIONESCUOLA"], ordine, an["CODICEISTITUTORIFERIMENTO"], an["DENOMINAZIONEISTITUTORIFERIMENTO"],
                an["INDIRIZZOSCUOLA"], an["CAPSCUOLA"], com[0], com[1], pv[0], pv[1], reg[0], reg[1],
                d["tot"], d["str"], round(pct_tot, 1),
                "sì" if anni_over else "no", "sì" if tot_over else "no",
                len(anni_over), ", ".join(map(str, anni_over)),
                *[pct_anni.get(y) for y in range(1, 6)],
            ])

    # ---- foglio province
    hp = ["Regione", "Codice Istat regione", "Provincia", "Codice Istat provincia"]
    for o, lab in [("Primaria", "Primarie"), ("Secondaria di primo grado", "Secondarie I grado")]:
        hp += [f"{lab} (n)", f"{lab} con almeno un anno >30% (n)", f"{lab} con almeno un anno >30% (%)",
               f"{lab} con totale >30% (n)", f"{lab} con totale >30% (%)"]
    rows_p = []
    for (reg, pv), c in sorted(prov.items(), key=lambda x: (x[0][0][1], x[0][1][0])):
        row = [reg[0], reg[1], pv[0], pv[1]]
        for o in ["Primaria", "Secondaria di primo grado"]:
            n = c[o + "|n"]
            row += [n, c[o + "|anno"], round(c[o + "|anno"] / n * 100, 1) if n else None,
                    c[o + "|tot"], round(c[o + "|tot"] / n * 100, 1) if n else None]
        rows_p.append(row)

    he = ["Codice scuola", "Denominazione scuola", "Ordine di scuola", "Codice istituto di riferimento",
          "Denominazione istituto di riferimento", "Indirizzo", "CAP", "Comune", "Codice Istat comune",
          "Provincia", "Codice Istat provincia", "Regione", "Codice Istat regione",
          "Alunni (n)", f"Alunni stranieri {etichetta} (n)", f"Alunni stranieri {etichetta} (%)",
          "Almeno un anno di corso >30%", "Totale scuola >30%", "Anni di corso >30% (n)", "Anni di corso >30%",
          "Stranieri 1° anno (%)", "Stranieri 2° anno (%)", "Stranieri 3° anno (%)", "Stranieri 4° anno (%)", "Stranieri 5° anno (%)"]
    elenco.sort(key=lambda r: (r[12], r[9], r[7], r[2], r[1]))

    text_cols_e = {"Codice scuola", "CAP", "Codice Istat comune", "Codice Istat provincia", "Codice Istat regione", "Anni di corso >30%"}
    wb = Workbook()
    write(wb.active, "Province", hp, rows_p, {"Codice Istat regione", "Codice Istat provincia"})
    write(wb.create_sheet(), "Scuole", he, elenco, text_cols_e)
    wb.save(BASE / out_name)

    # riepilogo nazionale
    tot = C.Counter()
    for c in prov.values(): tot.update(c)
    print(out_name, "| scuole in elenco:", len(elenco))
    for o in ["Primaria", "Secondaria di primo grado"]:
        print(f"  {o}: {tot[o+'|n']} plessi, anno>30%: {tot[o+'|anno']} ({tot[o+'|anno']/tot[o+'|n']*100:.1f}%), totale>30%: {tot[o+'|tot']} ({tot[o+'|tot']/tot[o+'|n']*100:.1f}%)")
    top = sorted(rows_p, key=lambda r: -r[6])[:5]
    print("  top primarie (anno):", [(r[2], r[6]) for r in top])
    top = sorted(rows_p, key=lambda r: -r[11])[:5]
    print("  top secondarie (anno):", [(r[2], r[11]) for r in top])

def fmt_for(vals):
    v = [x for x in vals if isinstance(x, (int, float))]
    if all(float(x).is_integer() for x in v): return "0"
    if all(round(x, 1) == x for x in v): return "0.0"
    return "0.00"

def write(ws, title, header, rows, text_cols):
    ws.title = title
    hf, df = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10)
    L, R = Alignment(horizontal="left"), Alignment(horizontal="right")
    for j, h in enumerate(header):
        col = [r[j] for r in rows]
        numeric = h not in text_cols and any(isinstance(x, (int, float)) for x in col)
        fmt = fmt_for(col) if numeric else "@"
        c = ws.cell(row=1, column=j + 1, value=h); c.font = hf; c.alignment = R if numeric else L
        for i, v in enumerate(col, start=2):
            if isinstance(v, float) and v.is_integer() and fmt == "0": v = int(v)
            c = ws.cell(row=i, column=j + 1, value=v)
            c.font = df; c.alignment = R if numeric else L; c.number_format = fmt
        width = max([len(str(h))] + [len(str(x)) for x in col[:3000] if x is not None])
        ws.column_dimensions[c.column_letter].width = min(max(width + 2, 8), 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

build("ALUNNICITTADINANZANONITALIANAPAESINONUE", "non UE", "Scuole_stranieri_oltre30_nonUE_2024-25.xlsx")
build("ALUNNICITTADINANZANONITALIANA", "UE e non UE", "Scuole_stranieri_oltre30_UE_e_nonUE_2024-25.xlsx")
print("comuni senza codice:", sorted(set(non_match)))
